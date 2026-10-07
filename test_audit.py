"""
Testes unitários e de estresse para o Motor de Auditoria do Treinamento de Integração BRVAL.
Cobre os requisitos da Seção 4.1 do SDD:
- Teste de Estresse de Datas (anos bissextos, precisão com divisor 365.25)
- Simulação matemática da tabela 2.2.1 do SDD
- Teste de Formato e Resiliência (arquivos corrompidos, PDFs, colunas ausentes)
- Validação com a planilha real 'Controle de Treinamento 03 anos.xlsx'
"""

import io
import unittest
from datetime import datetime
import pandas as pd

from audit_engine import (
    process_training_data,
    read_training_file,
    generate_export_csv,
    InvalidFileFormatError,
    MissingColumnsError,
    AuditException
)


class TestBRVALAuditEngine(unittest.TestCase):

    def test_sdd_simulation_matrix(self):
        """
        Valida a tabela de simulação matemática especificada na seção 2.2.1 do SDD:
        Func A: Admissão 01/01/2020, 1º Treinamento 01/01/2023 -> Data Base 01/01/2023 -> Status: Precisa
        Func B: Admissão 15/05/2023, 1º Treinamento Nulo       -> Data Base 15/05/2023 -> Status: Precisa
        Func C: Admissão 10/10/2024, 1º Treinamento Nulo       -> Data Base 10/10/2024 -> Status: Não Precisa
        """
        # Data de referência fixada em 06/10/2026 (data atual do documento)
        ref_date = datetime(2026, 10, 6)

        data = {
            "Empresa ": ["BRVAL MATRIZ", "BRVAL SERVICE", "BRVAL TRANSFORMADORES"],
            "Nome do Colaborador ": ["Func. A", "Func. B", "Func. C"],
            "Admissão ": ["2020-01-01", "2023-05-15", "2024-10-10"],
            "Demissão": [None, None, None],
            "1º Treinamento": ["2023-01-01", None, None],
        }
        df = pd.DataFrame(data)

        df_result, kpis = process_training_data(df, reference_date=ref_date)

        # Func. A: Data Base 01/01/2023. De 01/01/2023 até 06/10/2026 = 1374 dias / 365.25 = 3.76 anos (> 3)
        func_a = df_result[df_result["Colaborador"] == "Func. A"].iloc[0]
        self.assertEqual(func_a["Data_Base"].strftime("%Y-%m-%d"), "2023-01-01")
        self.assertEqual(func_a["Origem_Data_Base"], "1º Treinamento")
        self.assertGreater(func_a["Tempo_Decorrido_Anos"], 3.0)
        self.assertEqual(func_a["Status"], "Precisa")

        # Func. B: Data Base 15/05/2023. De 15/05/2023 até 06/10/2026 = 1240 dias / 365.25 = 3.39 anos (> 3)
        func_b = df_result[df_result["Colaborador"] == "Func. B"].iloc[0]
        self.assertEqual(func_b["Data_Base"].strftime("%Y-%m-%d"), "2023-05-15")
        self.assertEqual(func_b["Origem_Data_Base"], "Admissão")
        self.assertGreater(func_b["Tempo_Decorrido_Anos"], 3.0)
        self.assertEqual(func_b["Status"], "Precisa")

        # Func. C: Data Base 10/10/2024. De 10/10/2024 até 06/10/2026 = 726 dias / 365.25 = 1.99 anos (<= 3)
        func_c = df_result[df_result["Colaborador"] == "Func. C"].iloc[0]
        self.assertEqual(func_c["Data_Base"].strftime("%Y-%m-%d"), "2024-10-10")
        self.assertEqual(func_c["Origem_Data_Base"], "Admissão")
        self.assertLessEqual(func_c["Tempo_Decorrido_Anos"], 3.0)
        self.assertEqual(func_c["Status"], "Não Precisa")

        self.assertEqual(kpis["total_precisam"], 2)
        self.assertEqual(kpis["total_conforme"], 1)

    def test_leap_year_stress_test(self):
        """
        Teste de Estresse de Anos Bissextos:
        Verifica se datas de anos bissextos (como 29/02/2020 e 29/02/2024)
        mantêm a acurácia matemática através do divisor 365.25.
        """
        ref_date = datetime(2024, 3, 1)

        # Exatamente 4 anos antes em ano bissexto: 29/02/2020
        # De 29/02/2020 até 01/03/2024 são 1462 dias / 365.25 = 4.002 anos -> 'Precisa'
        data = {
            "Empresa": ["BRVAL TESTE"],
            "Colaborador": ["Colab Bissexto"],
            "Admissao": ["2020-02-29"],
            "Demissao": [None],
            "1º Treinamento": [None]
        }
        df = pd.DataFrame(data)
        df_result, kpis = process_training_data(df, reference_date=ref_date)
        colab = df_result.iloc[0]
        self.assertAlmostEqual(colab["Tempo_Decorrido_Anos"], 4.0, delta=0.05)
        self.assertEqual(colab["Status"], "Precisa")

    def test_dismissed_employees_are_ignored(self):
        """
        Verifica se colaboradores demitidos (com data de demissão preenchida)
        são estritamente ignorados conforme o requisito funcional 1.2 e seção 3.2.
        """
        data = {
            "Empresa": ["BRVAL A", "BRVAL B"],
            "Colaborador": ["Demitido", "Ativo"],
            "Admissao": ["2015-01-01", "2025-01-01"],
            "Demissao": ["2022-01-01", None],
            "1º Treinamento": [None, None]
        }
        df = pd.DataFrame(data)
        df_result, kpis = process_training_data(df)

        self.assertEqual(len(df_result), 1)
        self.assertEqual(df_result.iloc[0]["Colaborador"], "Ativo")
        self.assertEqual(kpis["total_demitidos"], 1)
        self.assertEqual(kpis["total_ativos"], 1)

    def test_corrupted_file_handling(self):
        """
        Teste de Formato (Seção 4.1): Inserir arquivos corrompidos ou PDFs
        garantindo erro amigável sem stack trace inesperado.
        """
        fake_pdf = io.BytesIO(b"%PDF-1.4 Fake PDF Content Not An Excel")
        setattr(fake_pdf, "name", "documento.pdf")

        with self.assertRaises(InvalidFileFormatError) as ctx:
            read_training_file(fake_pdf)
        self.assertIn("Formato de arquivo não suportado", str(ctx.exception))

        # Teste com arquivo corrompido nomeado como xlsx
        corrupt_xlsx = io.BytesIO(b"Not a real zip archive")
        setattr(corrupt_xlsx, "name", "planilha_corrompida.xlsx")

        with self.assertRaises(InvalidFileFormatError) as ctx:
            read_training_file(corrupt_xlsx)
        self.assertIn("Não foi possível abrir o arquivo", str(ctx.exception))

    def test_missing_required_columns(self):
        """
        Verifica se a ausência de colunas obrigatórias dispara MissingColumnsError amigável.
        """
        data_sem_admissao = {
            "Empresa": ["BRVAL A"],
            "Nome do Colaborador": ["Teste"],
        }
        df = pd.DataFrame(data_sem_admissao)
        with self.assertRaises(MissingColumnsError) as ctx:
            process_training_data(df)
        self.assertIn("Admissão", str(ctx.exception))

    def test_real_excel_file_audit(self):
        """
        Executa a auditoria na planilha real 'Controle de Treinamento 03 anos.xlsx'
        e confere consistência e geração de CSV.
        """
        df = read_training_file("Controle de Treinamento 03 anos.xlsx")
        self.assertGreater(len(df), 300)

        df_result, kpis = process_training_data(df)
        self.assertEqual(kpis["total_registros_planilha"], 367)
        self.assertEqual(kpis["total_demitidos"], 231)
        self.assertEqual(kpis["total_ativos"], 136)
        self.assertGreater(kpis["total_precisam"], 0)

        # Testa exportação CSV
        csv_bytes = generate_export_csv(df_result[df_result["Status"] == "Precisa"])
        self.assertTrue(csv_bytes.startswith(b"\xef\xbb\xbf"))  # UTF-8 BOM

    def test_executive_excel_generation(self):
        """
        Valida a geração de planilha Excel (.xlsx) corporativa BRVAL:
        Verifica integridade do binário, presença das 3 abas essenciais e estrutura.
        """
        import openpyxl
        from audit_engine import generate_export_excel

        df = read_training_file("Controle de Treinamento 03 anos.xlsx")
        df_result, kpis = process_training_data(df)
        excel_bytes = generate_export_excel(df_result, kpis)

        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 5000)

        # Carrega o workbook a partir dos bytes para conferir abas
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        self.assertIn("Ação Imediata - Vencidos", sheet_names)
        self.assertIn("Base Ativos Completa", sheet_names)
        self.assertIn("Quadro Resumo Filiais", sheet_names)

        # Confere que a aba de vencidos contém cabeçalho e dados
        ws_vencidos = wb["Ação Imediata - Vencidos"]
        self.assertIn("BRVAL", ws_vencidos["A1"].value)
        self.assertGreater(ws_vencidos.max_row, 4)

    def test_data_quality_auditing(self):
        """
        Verifica se o motor de Data Hygiene detecta anomalias cadastrais
        como admissão no futuro e treinamento anterior à admissão.
        """
        from audit_engine import check_data_quality

        data_com_anomalias = {
            "Empresa": ["BRVAL TESTE", "BRVAL TESTE 2"],
            "Nome do Colaborador": ["Colab Futuro", "Colab Trein Antes"],
            "Admissao": ["2030-01-01", "2024-05-10"],  # Futuro e normal
            "Demissao": [None, None],
            "1º Treinamento": [None, "2020-01-01"]     # Treinamento antes da admissão
        }
        df_inconsistente = pd.DataFrame(data_com_anomalias)
        ref_date = datetime(2026, 10, 6)
        issues = check_data_quality(df_inconsistente, reference_date=ref_date)

        self.assertEqual(len(issues), 2)
        motivos = [i["motivo"] for i in issues]
        self.assertTrue(any("está no futuro" in m for m in motivos))
        self.assertTrue(any("é anterior à Admissão" in m for m in motivos))

    def test_convocation_templates_generation(self):
        """
        Valida a geração de convocação formal individual e em lote.
        """
        from audit_engine import generate_convocation_email, generate_batch_convocation

        email_text = generate_convocation_email(
            colaborador_nome="Carlos Silva",
            empresa="BRVAL MATRIZ",
            tempo_decorrido=3.85,
            data_base_str="01/01/2023",
            origem_data_base="1º Treinamento"
        )
        self.assertIn("Carlos Silva", email_text)
        self.assertIn("BRVAL MATRIZ", email_text)
        self.assertIn("S-2245", email_text)
        self.assertIn("NR-1", email_text)

        batch_text = generate_batch_convocation(
            empresa="BRVAL SERVICE",
            colaboradores=[
                {"Colaborador": "Ana Costa", "Tempo_Decorrido_Anos": 4.12, "Data_Base_Fmt": "10/05/2022"},
                {"Colaborador": "Marcos Dias", "Tempo_Decorrido_Anos": 3.40, "Data_Base_Fmt": "15/08/2023"}
            ]
        )
        self.assertIn("Total de Convocados: 2", batch_text)
        self.assertIn("Ana Costa", batch_text)
        self.assertIn("Marcos Dias", batch_text)


if __name__ == "__main__":
    unittest.main()

