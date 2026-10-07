"""
BRVAL Hub de Treinamentos - Motor de Auditoria e Regras de Negócio (Framework LACE)
Módulo de processamento de dados e cálculo de elegibilidade para o Treinamento de Integração.
Atende aos requisitos de privacidade (LGPD): processamento estritamente em memória volátil.
"""

import io
import unicodedata
from datetime import datetime
from typing import Dict, Any, Tuple, Optional
import pandas as pd


class AuditException(Exception):
    """Exceção base amigável para regras de auditoria."""
    pass


class InvalidFileFormatError(AuditException):
    """Lançada quando o formato do arquivo não é suportado ou está corrompido."""
    pass


class MissingColumnsError(AuditException):
    """Lançada quando a planilha não contém as colunas mínimas exigidas."""
    pass


def normalize_string(text: str) -> str:
    """Remove acentos, espaços extras e converte para minúsculas para comparação flexível."""
    if not isinstance(text, str):
        text = str(text) if text is not None else ""
    nfkd = unicodedata.normalize("NFKD", text)
    ascii_text = nfkd.encode("ASCII", "ignore").decode("utf-8")
    return ascii_text.strip().lower()


def find_column_match(columns: list, patterns: list) -> Optional[str]:
    """Localiza o nome original da coluna baseado em padrões flexíveis."""
    for col in columns:
        norm = normalize_string(col)
        for pat in patterns:
            if pat in norm:
                return col
    return None


def read_training_file(file_input: Any) -> pd.DataFrame:
    """
    Lê o arquivo de treinamento com tratamento amigável de erros de formato.
    Suporta .xlsx, .xls e .csv.
    """
    try:
        if isinstance(file_input, str):
            if file_input.endswith(".csv"):
                return pd.read_csv(file_input)
            return pd.read_excel(file_input)
        elif hasattr(file_input, "name"):
            filename = file_input.name.lower()
            if filename.endswith(".csv"):
                return pd.read_csv(file_input)
            elif filename.endswith((".xlsx", ".xls")):
                return pd.read_excel(file_input)
            else:
                raise InvalidFileFormatError(
                    f"Formato de arquivo não suportado: '{file_input.name}'. "
                    "Por favor, envie uma planilha em formato Excel (.xlsx ou .xls) ou CSV."
                )
        elif isinstance(file_input, (bytes, io.BytesIO)):
            # Tenta ler como Excel
            try:
                return pd.read_excel(file_input)
            except Exception:
                if isinstance(file_input, bytes):
                    file_input = io.BytesIO(file_input)
                file_input.seek(0)
                return pd.read_csv(file_input)
        else:
            raise InvalidFileFormatError("Entrada de arquivo inválida.")
    except InvalidFileFormatError:
        raise
    except Exception as e:
        raise InvalidFileFormatError(
            "Não foi possível abrir o arquivo fornecido. O arquivo pode estar corrompido, "
            "vazio ou em um formato incompatível (como PDF renomeado). "
            f"Detalhe técnico: {str(e)}"
        )


def process_training_data(
    df_raw: pd.DataFrame,
    reference_date: Optional[datetime] = None
) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    """
    Executa o motor de regras de negócio da BRVAL:
    1. Localiza e valida colunas essenciais.
    2. Filtra e descarta colaboradores demitidos (df['Demissão'].isna()).
    3. Define a Data Base: '1º Treinamento' se preenchido, senão 'Admissão'.
    4. Calcula o tempo decorrido em anos: (Data Atual - Data Base) / 365.25.
    5. Aplica a matriz de decisão: Status = 'Precisa' se Tempo > 3, senão 'Não Precisa'.
    6. Retorna o DataFrame processado e indicadores agregados.
    """
    if df_raw is None or df_raw.empty:
        raise AuditException("A planilha enviada está vazia ou não contém linhas de dados.")

    columns = df_raw.columns.tolist()

    # Mapeamento inteligente de colunas
    col_empresa = find_column_match(columns, ["empresa", "filial", "unidade"])
    col_nome = find_column_match(columns, ["nome do colaborador", "colaborador", "nome", "funcionario"])
    col_admissao = find_column_match(columns, ["admissao", "data admissao", "admitido"])
    col_demissao = find_column_match(columns, ["demissao", "data demissao", "rescisao"])
    col_treinamento = find_column_match(columns, ["1o treinamento", "1 treinamento", "primeiro treinamento", "treinamento"])

    missing = []
    if not col_empresa:
        missing.append("Empresa")
    if not col_nome:
        missing.append("Nome do Colaborador")
    if not col_admissao:
        missing.append("Admissão")

    if missing:
        raise MissingColumnsError(
            f"A planilha enviada não contém as seguintes colunas obrigatórias: {', '.join(missing)}. "
            "Por favor, verifique o cabeçalho do arquivo 'Controle de Treinamento 03 anos.xlsx'."
        )

    total_registros = len(df_raw)

    # Criação do DataFrame limpo e padronizado
    clean_df = pd.DataFrame()
    clean_df["Empresa"] = df_raw[col_empresa].astype(str).str.strip()
    clean_df["Colaborador"] = df_raw[col_nome].astype(str).str.strip()
    clean_df["Admissao"] = pd.to_datetime(df_raw[col_admissao], errors="coerce")

    if col_demissao:
        clean_df["Demissao"] = pd.to_datetime(df_raw[col_demissao], errors="coerce")
    else:
        clean_df["Demissao"] = pd.NaT

    if col_treinamento:
        clean_df["Primeiro_Treinamento"] = pd.to_datetime(df_raw[col_treinamento], errors="coerce")
    else:
        clean_df["Primeiro_Treinamento"] = pd.NaT

    # Filtro: Ignorar colaboradores demitidos (df['Demissão'].isna())
    demitidos_count = clean_df["Demissao"].notna().sum()
    ativos = clean_df[clean_df["Demissao"].isna()].copy()
    total_ativos = len(ativos)

    if total_ativos == 0:
        raise AuditException("Nenhum colaborador ativo encontrado na planilha (todos constam com data de demissão).")

    # Definição da Data Base:
    # Se houver data na coluna "1º Treinamento", ela será a Data Base. Caso contrário, assume-se a "Admissão".
    ativos["Data_Base"] = ativos["Primeiro_Treinamento"].fillna(ativos["Admissao"])
    ativos["Origem_Data_Base"] = ativos["Primeiro_Treinamento"].notna().map({True: "1º Treinamento", False: "Admissão"})

    # Data de Referência (Data Atual ou especificada para testes)
    data_ref = pd.Timestamp(reference_date if reference_date is not None else datetime.now()).normalize()

    # Fórmula Matemática de Tempo:
    # Tempo = (Data Atual - Data Base) / 365.25
    # Garantir que a diferença em dias considere anos bissextos
    delta_dias = (data_ref - ativos["Data_Base"]).dt.days
    ativos["Tempo_Decorrido_Anos"] = (delta_dias / 365.25).round(2)

    # Matriz de Decisão Lógica:
    # Status = "Precisa", se Tempo > 3; "Não Precisa", se Tempo <= 3
    ativos["Status"] = ativos["Tempo_Decorrido_Anos"].apply(
        lambda t: "Precisa" if (pd.notna(t) and t > 3.0) else "Não Precisa"
    )

    # Colunas formatadas para exibição amigável
    ativos["Admissao_Fmt"] = ativos["Admissao"].dt.strftime("%d/%m/%Y").fillna("-")
    ativos["Primeiro_Treinamento_Fmt"] = ativos["Primeiro_Treinamento"].dt.strftime("%d/%m/%Y").fillna("-")
    ativos["Data_Base_Fmt"] = ativos["Data_Base"].dt.strftime("%d/%m/%Y").fillna("-")

    # Data Limite de Renovação (Data Base + 3 anos exatos ou 1095.75 dias)
    ativos["Data_Limite_Renovacao"] = ativos["Data_Base"] + pd.to_timedelta(3 * 365.25, unit="D")
    ativos["Data_Limite_Fmt"] = ativos["Data_Limite_Renovacao"].dt.strftime("%d/%m/%Y").fillna("-")
    ativos["Dias_Para_Vencer"] = (ativos["Data_Limite_Renovacao"] - data_ref).dt.days

    # Classificação de Nível de Urgência / Risco Trabalhista:
    # - Crítico: > 5 anos de defasagem (risco gravíssimo em fiscalização)
    # - Alto: 4 a 5 anos
    # - Moderado: 3 a 4 anos
    # - Atenção Preventiva: vence nos próximos 90 dias (2.75 a 3.0 anos)
    # - Conforme: <= 2.75 anos
    def classify_urgency(tempo: float) -> str:
        if pd.isna(tempo):
            return "Não Identificado"
        if tempo > 5.0:
            return "Crítico (> 5 anos)"
        if tempo > 4.0:
            return "Alto (4 a 5 anos)"
        if tempo > 3.0:
            return "Moderado (3 a 4 anos)"
        if tempo > 2.75:
            return "Atenção (Vence em 90 dias)"
        return "Conforme Regular"

    ativos["Nivel_Risco"] = ativos["Tempo_Decorrido_Anos"].apply(classify_urgency)

    # Indicadores
    total_precisam = (ativos["Status"] == "Precisa").sum()
    total_conforme = (ativos["Status"] == "Não Precisa").sum()
    total_criticos = (ativos["Nivel_Risco"] == "Crítico (> 5 anos)").sum()
    total_atencao_90d = (ativos["Nivel_Risco"] == "Atenção (Vence em 90 dias)").sum()
    taxa_conformidade = round((total_conforme / total_ativos) * 100, 1) if total_ativos > 0 else 0.0
    empresas_unicas = sorted(ativos["Empresa"].unique().tolist())

    kpis = {
        "total_registros_planilha": total_registros,
        "total_demitidos": int(demitidos_count),
        "total_ativos": int(total_ativos),
        "total_precisam": int(total_precisam),
        "total_conforme": int(total_conforme),
        "total_criticos": int(total_criticos),
        "total_atencao_90d": int(total_atencao_90d),
        "taxa_conformidade": taxa_conformidade,
        "total_empresas": len(empresas_unicas),
        "empresas": empresas_unicas,
        "data_referencia": data_ref.strftime("%d/%m/%Y"),
        "data_referencia_obj": data_ref,
    }

    return ativos, kpis


def check_data_quality(df_raw: pd.DataFrame, reference_date: Optional[datetime] = None) -> list:
    """
    Realiza auditoria de higienização cadastral (Data Hygiene):
    Detecta anomalias que prejudicam cruzamentos no eSocial (S-2245).
    """
    issues = []
    if df_raw is None or df_raw.empty:
        return issues

    columns = df_raw.columns.tolist()
    col_empresa = find_column_match(columns, ["empresa", "filial", "unidade"])
    col_nome = find_column_match(columns, ["nome do colaborador", "colaborador", "nome", "funcionario"])
    col_admissao = find_column_match(columns, ["admissao", "data admissao", "admitido"])
    col_demissao = find_column_match(columns, ["demissao", "data demissao", "rescisao"])
    col_treinamento = find_column_match(columns, ["1o treinamento", "1 treinamento", "primeiro treinamento", "treinamento"])

    data_ref = pd.Timestamp(reference_date if reference_date is not None else datetime.now()).normalize()

    for idx, row in df_raw.iterrows():
        linha_num = idx + 2  # Excel 1-indexed com cabeçalho
        colab = str(row.get(col_nome, "")).strip() if col_nome else f"Linha {linha_num}"
        empresa = str(row.get(col_empresa, "")).strip() if col_empresa else "Não informada"

        adm_val = pd.to_datetime(row.get(col_admissao), errors="coerce") if col_admissao else pd.NaT
        dem_val = pd.to_datetime(row.get(col_demissao), errors="coerce") if col_demissao else pd.NaT
        trein_val = pd.to_datetime(row.get(col_treinamento), errors="coerce") if col_treinamento else pd.NaT

        # 1. Admissão no futuro
        if pd.notna(adm_val) and adm_val > data_ref:
            issues.append({
                "gravidade": "Atenção",
                "colaborador": colab,
                "empresa": empresa,
                "motivo": f"Data de Admissão ({adm_val.strftime('%d/%m/%Y')}) está no futuro.",
                "linha": linha_num
            })

        # 2. 1º Treinamento anterior à Admissão
        if pd.notna(adm_val) and pd.notna(trein_val) and trein_val < adm_val:
            issues.append({
                "gravidade": "Inconsistência",
                "colaborador": colab,
                "empresa": empresa,
                "motivo": f"1º Treinamento ({trein_val.strftime('%d/%m/%Y')}) é anterior à Admissão ({adm_val.strftime('%d/%m/%Y')}).",
                "linha": linha_num
            })

        # 3. Demissão anterior à Admissão
        if pd.notna(adm_val) and pd.notna(dem_val) and dem_val < adm_val:
            issues.append({
                "gravidade": "Erro Cadastral",
                "colaborador": colab,
                "empresa": empresa,
                "motivo": f"Data de Demissão ({dem_val.strftime('%d/%m/%Y')}) anterior à Admissão.",
                "linha": linha_num
            })

    return issues


def generate_export_csv(df_to_export: pd.DataFrame) -> bytes:
    """
    Gera um CSV formatado em UTF-8 com BOM (compatível com Excel em português)
    com as colunas e formatos ideais para acionamento das lideranças de DP.
    """
    export_df = pd.DataFrame()
    export_df["Empresa"] = df_to_export["Empresa"]
    export_df["Nome do Colaborador"] = df_to_export["Colaborador"]
    export_df["Data de Admissão"] = df_to_export["Admissao_Fmt"]
    export_df["1º Treinamento"] = df_to_export["Primeiro_Treinamento_Fmt"]
    export_df["Data Base Utilizada"] = df_to_export["Data_Base_Fmt"]
    export_df["Origem Data Base"] = df_to_export["Origem_Data_Base"]
    export_df["Tempo Decorrido (Anos)"] = df_to_export["Tempo_Decorrido_Anos"]
    export_df["Data Limite Renovação"] = df_to_export.get("Data_Limite_Fmt", "-")
    export_df["Nível de Risco"] = df_to_export.get("Nivel_Risco", "-")
    export_df["Status Auditoria"] = df_to_export["Status"]

    csv_buffer = io.StringIO()
    export_df.to_csv(csv_buffer, index=False, sep=";", encoding="utf-8")
    # Adiciona UTF-8 BOM para garantir que acentos abram perfeitamente no Excel
    return "\ufeff".encode("utf-8") + csv_buffer.getvalue().encode("utf-8")


def generate_export_excel(
    df_ativos: pd.DataFrame,
    kpis: Dict[str, Any],
    empresa_filtro: Optional[str] = None
) -> bytes:
    """
    Gera uma planilha executiva em Excel (.xlsx) altamente estilizada
    com a identidade visual corporativa da BRVAL (Verde-água e Amarelo).
    Inclui 3 abas estruturadas:
    1. Ação Imediata (Vencidos): Colaboradores prioritários para agendamento.
    2. Base Auditada Completa: Todos os colaboradores ativos com status e risco.
    3. Resumo por Filial: Indicadores consolidados por unidade de negócio.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # Remove sheet padrão inicial
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Paleta BRVAL
    teal_fill = PatternFill(start_color="57A695", end_color="57A695", fill_type="solid")
    dark_fill = PatternFill(start_color="151C28", end_color="151C28", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFCC29", end_color="FFCC29", fill_type="solid")
    light_yellow_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    light_teal_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    red_alert_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="151C28")
    font_regular = Font(name="Calibri", size=10, color="000000")
    font_title = Font(name="Calibri", size=14, bold=True, color="57A695")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="555555")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    # -------------------------------------------------------------
    # ABA 1: AÇÃO IMEDIATA (ELEGÍVEIS / VENCIDOS)
    # -------------------------------------------------------------
    ws_vencidos = wb.create_sheet(title="Ação Imediata - Vencidos")
    ws_vencidos.views.sheetView[0].showGridLines = True

    # Cabeçalho do Relatório
    ws_vencidos["A1"] = "BRVAL - Hub de Treinamentos (Relatório de Ação Imediata)"
    ws_vencidos["A1"].font = font_title
    subtitulo = f"Auditoria de Elegibilidade: Treinamento de Integração (Ciclo de 3 Anos) | Data Base: {kpis.get('data_referencia', '')}"
    if empresa_filtro and empresa_filtro != "Todas as Empresas":
        subtitulo += f" | Filtro: {empresa_filtro}"
    ws_vencidos["A2"] = subtitulo
    ws_vencidos["A2"].font = font_subtitle

    headers_vencidos = [
        "Empresa",
        "Nome do Colaborador",
        "Admissão",
        "1º Treinamento",
        "Data Base Utilizada",
        "Origem Data Base",
        "Tempo Decorrido (Anos)",
        "Limite Renovação",
        "Nível de Urgência",
        "Status"
    ]

    # Preenche cabeçalhos na linha 4
    for col_idx, header in enumerate(headers_vencidos, start=1):
        cell = ws_vencidos.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = teal_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_vencidos.row_dimensions[4].height = 24

    # Dados dos Vencidos
    df_vencidos = df_ativos[df_ativos["Status"] == "Precisa"].sort_values(by="Tempo_Decorrido_Anos", ascending=False)
    if empresa_filtro and empresa_filtro != "Todas as Empresas":
        df_vencidos = df_vencidos[df_vencidos["Empresa"] == empresa_filtro]

    row_num = 5
    for _, row in df_vencidos.iterrows():
        ws_vencidos.cell(row=row_num, column=1, value=row["Empresa"]).alignment = Alignment(horizontal="left")
        ws_vencidos.cell(row=row_num, column=2, value=row["Colaborador"]).alignment = Alignment(horizontal="left")
        ws_vencidos.cell(row=row_num, column=3, value=row["Admissao_Fmt"]).alignment = Alignment(horizontal="center")
        ws_vencidos.cell(row=row_num, column=4, value=row["Primeiro_Treinamento_Fmt"]).alignment = Alignment(horizontal="center")
        ws_vencidos.cell(row=row_num, column=5, value=row["Data_Base_Fmt"]).alignment = Alignment(horizontal="center")
        ws_vencidos.cell(row=row_num, column=6, value=row["Origem_Data_Base"]).alignment = Alignment(horizontal="center")
        
        c_tempo = ws_vencidos.cell(row=row_num, column=7, value=row["Tempo_Decorrido_Anos"])
        c_tempo.alignment = Alignment(horizontal="right")
        c_tempo.number_format = "0.00"

        ws_vencidos.cell(row=row_num, column=8, value=row.get("Data_Limite_Fmt", "-")).alignment = Alignment(horizontal="center")
        
        c_urgencia = ws_vencidos.cell(row=row_num, column=9, value=row.get("Nivel_Risco", "-"))
        c_urgencia.alignment = Alignment(horizontal="center")
        if "Crítico" in str(row.get("Nivel_Risco", "")):
            c_urgencia.fill = red_alert_fill
        else:
            c_urgencia.fill = light_yellow_fill

        c_status = ws_vencidos.cell(row=row_num, column=10, value="Precisa")
        c_status.alignment = Alignment(horizontal="center")
        c_status.font = font_bold
        c_status.fill = light_yellow_fill

        for col_idx in range(1, 11):
            ws_vencidos.cell(row=row_num, column=col_idx).border = thin_border
            if ws_vencidos.cell(row=row_num, column=col_idx).font == font_header:
                pass
            elif col_idx != 10:
                ws_vencidos.cell(row=row_num, column=col_idx).font = font_regular

        row_num += 1

    # Autoajuste de largura
    for col in ws_vencidos.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_vencidos.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws_vencidos.auto_filter.ref = f"A4:J{max(row_num - 1, 4)}"

    # -------------------------------------------------------------
    # ABA 2: BASE AUDITADA COMPLETA (TODOS OS ATIVOS)
    # -------------------------------------------------------------
    ws_all = wb.create_sheet(title="Base Ativos Completa")
    ws_all.views.sheetView[0].showGridLines = True

    ws_all["A1"] = "BRVAL - Base Completa de Colaboradores Ativos Auditados"
    ws_all["A1"].font = font_title
    ws_all["A2"] = f"Total Ativos: {kpis.get('total_ativos', 0)} | Em Conformidade: {kpis.get('total_conforme', 0)} | Vencidos: {kpis.get('total_precisam', 0)}"
    ws_all["A2"].font = font_subtitle

    for col_idx, header in enumerate(headers_vencidos, start=1):
        cell = ws_all.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = dark_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_all.row_dimensions[4].height = 24

    df_base = df_ativos.copy()
    if empresa_filtro and empresa_filtro != "Todas as Empresas":
        df_base = df_base[df_base["Empresa"] == empresa_filtro]

    row_num = 5
    for _, row in df_base.iterrows():
        ws_all.cell(row=row_num, column=1, value=row["Empresa"]).alignment = Alignment(horizontal="left")
        ws_all.cell(row=row_num, column=2, value=row["Colaborador"]).alignment = Alignment(horizontal="left")
        ws_all.cell(row=row_num, column=3, value=row["Admissao_Fmt"]).alignment = Alignment(horizontal="center")
        ws_all.cell(row=row_num, column=4, value=row["Primeiro_Treinamento_Fmt"]).alignment = Alignment(horizontal="center")
        ws_all.cell(row=row_num, column=5, value=row["Data_Base_Fmt"]).alignment = Alignment(horizontal="center")
        ws_all.cell(row=row_num, column=6, value=row["Origem_Data_Base"]).alignment = Alignment(horizontal="center")
        
        c_tempo = ws_all.cell(row=row_num, column=7, value=row["Tempo_Decorrido_Anos"])
        c_tempo.alignment = Alignment(horizontal="right")
        c_tempo.number_format = "0.00"

        ws_all.cell(row=row_num, column=8, value=row.get("Data_Limite_Fmt", "-")).alignment = Alignment(horizontal="center")
        
        c_urgencia = ws_all.cell(row=row_num, column=9, value=row.get("Nivel_Risco", "-"))
        c_urgencia.alignment = Alignment(horizontal="center")

        c_status = ws_all.cell(row=row_num, column=10, value=row["Status"])
        c_status.alignment = Alignment(horizontal="center")
        c_status.font = font_bold
        if row["Status"] == "Precisa":
            c_status.fill = light_yellow_fill
            if "Crítico" in str(row.get("Nivel_Risco", "")):
                c_urgencia.fill = red_alert_fill
            else:
                c_urgencia.fill = light_yellow_fill
        else:
            c_status.fill = light_teal_fill
            if "Atenção" in str(row.get("Nivel_Risco", "")):
                c_urgencia.fill = light_yellow_fill

        for col_idx in range(1, 11):
            ws_all.cell(row=row_num, column=col_idx).border = thin_border
            if col_idx != 10:
                ws_all.cell(row=row_num, column=col_idx).font = font_regular

        row_num += 1

    for col in ws_all.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_all.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws_all.auto_filter.ref = f"A4:J{max(row_num - 1, 4)}"

    # -------------------------------------------------------------
    # ABA 3: RESUMO POR FILIAL
    # -------------------------------------------------------------
    ws_resumo = wb.create_sheet(title="Quadro Resumo Filiais")
    ws_resumo.views.sheetView[0].showGridLines = True

    ws_resumo["A1"] = "BRVAL - Quadro Comparativo de Riscos por Filial"
    ws_resumo["A1"].font = font_title
    ws_resumo["A2"] = "Visão executiva agregada para reporte à Diretoria e Lideranças"
    ws_resumo["A2"].font = font_subtitle

    headers_resumo = ["Filial / Empresa", "Ativos Auditados", "Treinamento Vencido", "Em Conformidade", "Taxa de Risco (%)"]
    for col_idx, header in enumerate(headers_resumo, start=1):
        cell = ws_resumo.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = teal_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_resumo.row_dimensions[4].height = 24

    resumo_df = df_ativos.groupby("Empresa").agg(
        total_ativos=("Colaborador", "count"),
        vencidos=("Status", lambda s: (s == "Precisa").sum()),
        conforme=("Status", lambda s: (s == "Não Precisa").sum())
    ).reset_index()
    resumo_df["taxa_risco"] = ((resumo_df["vencidos"] / resumo_df["total_ativos"]) * 100).round(1)
    resumo_df = resumo_df.sort_values(by="vencidos", ascending=False)

    row_num = 5
    for _, r in resumo_df.iterrows():
        ws_resumo.cell(row=row_num, column=1, value=r["Empresa"]).alignment = Alignment(horizontal="left")
        ws_resumo.cell(row=row_num, column=2, value=int(r["total_ativos"])).alignment = Alignment(horizontal="right")
        
        c_venc = ws_resumo.cell(row=row_num, column=3, value=int(r["vencidos"]))
        c_venc.alignment = Alignment(horizontal="right")
        if r["vencidos"] > 0:
            c_venc.font = font_bold
            c_venc.fill = light_yellow_fill
        
        ws_resumo.cell(row=row_num, column=4, value=int(r["conforme"])).alignment = Alignment(horizontal="right")
        
        c_taxa = ws_resumo.cell(row=row_num, column=5, value=f"{r['taxa_risco']}%")
        c_taxa.alignment = Alignment(horizontal="right")
        if r["taxa_risco"] > 30.0:
            c_taxa.fill = red_alert_fill
            c_taxa.font = font_bold

        for col_idx in range(1, 6):
            ws_resumo.cell(row=row_num, column=col_idx).border = thin_border
            if ws_resumo.cell(row=row_num, column=col_idx).font != font_bold:
                ws_resumo.cell(row=row_num, column=col_idx).font = font_regular

        row_num += 1

    for col in ws_resumo.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_resumo.column_dimensions[col_letter].width = max(max_len + 4, 15)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer.getvalue()


def generate_convocation_email(
    colaborador_nome: str,
    empresa: str,
    tempo_decorrido: float,
    data_base_str: str,
    origem_data_base: str
) -> str:
    """
    Gera modelo oficial de convocação individual para Treinamento de Integração BRVAL.
    """
    return f"""Prezado(a) Gestor(a) e Colaborador(a) {colaborador_nome},

Assunto: [CONVOCAÇÃO OFICIAL] Reciclagem Obrigatória - Treinamento de Integração BRVAL ({empresa})

Informamos que, em conformidade com as Normas Regulamentadoras (NR-1 e NRs correlatas) e com a política interna de Segurança e Compliance da BRVAL, identificamos a necessidade de reciclagem periódica do Treinamento de Integração para o colaborador abaixo:

• Colaborador: {colaborador_nome}
• Empresa / Unidade: {empresa}
• Referência Base ({origem_data_base}): {data_base_str}
• Tempo Decorrido: {tempo_decorrido:.2f} anos (Prazo regulamentar de 3 anos excedido)

ORIENTAÇÕES IMPORTANTES PARA O DEPARTAMENTO PESSOAL E GESTÃO:
1. Jornada de Trabalho (CLT Art. 4º): O treinamento deve ser ministrado obrigatoriamente durante o horário normal de trabalho.
2. Comprovação Legal e eSocial (Evento S-2245): A presença é de caráter obrigatório, sendo indispensável a assinatura da lista de presença física/digital para baixa da pendência fiscal.
3. Agendamento: Solicitamos contato com o SESMT / RH da BRVAL nas próximas 48 horas para inclusão na próxima turma disponível.

Atenciosamente,

Departamento Pessoal & Recursos Humanos
BRVAL Hub de Treinamentos
"""


def generate_batch_convocation(
    empresa: str,
    colaboradores: list
) -> str:
    """
    Gera comunicado em lote para o gestor da unidade e SESMT com a relação dos convocados.
    """
    linhas_colab = []
    for idx, c in enumerate(colaboradores, start=1):
        nome = c.get("Colaborador", "N/D")
        tempo = c.get("Tempo_Decorrido_Anos", 0)
        base = c.get("Data_Base_Fmt", "-")
        linhas_colab.append(f"{idx}. {nome} | Base: {base} | Tempo: {tempo:.2f} anos")

    lista_formatada = "\n".join(linhas_colab)

    return f"""MEMORANDO CIRCULAR DEPARTAMENTO PESSOAL BRVAL

Para: Liderança Operacional & SESMT - {empresa}
De: Departamento Pessoal & Compliance Trabalhista BRVAL
Assunto: Convocação para Formação de Turma - Treinamento de Integração Obrigatório
Total de Convocados: {len(colaboradores)} colaborador(es)

Prezada Liderança,

Em auditoria periódica de conformidade legal e trabalhista, constatamos que os seguintes colaboradores lotados na unidade {empresa} atingiram o ciclo de 3 anos desde a sua última integração/admissão:

{lista_formatada}

AÇÕES NECESSÁRIAS:
1. Liberar os colaboradores listados para realização do treinamento na grade definida pelo SESMT.
2. Assegurar que o treinamento ocorra estritamente dentro da jornada de trabalho regular, evitando acúmulo de horas extras.
3. Coletar e arquivar a ata de presença com as respectivas assinaturas para transmissão ao eSocial (Evento S-2245).

Contamos com o seu apoio para manter a segurança jurídica e operacional da BRVAL.

Departamento Pessoal BRVAL
"""

